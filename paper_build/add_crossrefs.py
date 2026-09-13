#!/usr/bin/env python
"""Turn plain-text references into real Word cross-reference fields.

The builder writes references as literal text ("Fig. 7", "Table IX", "[47]").
Word treats those as ordinary characters: they do not hyperlink, they do not
appear in the navigation pane, and Update Fields does nothing to them. Editors
and typesetters generally expect genuine cross-references.

This post-processes a built DOCX in two passes:

  1. bookmark the *label token only* in each target -- the "7" in "FIGURE 7.",
     the "IX" in "TABLE IX", the "2" in "Algorithm 2", the "[47]" in a
     bibliography entry;
  2. replace each in-prose reference with the literal label word plus a REF
     field carrying the \\h switch, so "Fig. 7" becomes the text "Fig. " followed
     by a field resolving to "7".

Bookmarking only the token is what makes Update Fields safe. A bookmark around
the whole caption paragraph looks correct until the first Ctrl+A F9, at which
point every reference expands into the entire caption and destroys the sentence.

Visible text is unchanged, so numbering stays under the control of
`paper_xref.py` and the build gates; what changes is that each reference becomes
clickable, updatable and navigable.

    python add_crossrefs.py [in.docx] [out.docx]

Run it AFTER build_access_paper.py. Rebuilding overwrites the DOCX and discards
the fields, so this is the last step before sending the file out.
"""
import copy
import os
import re
import sys

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

FIG_BM = "_Ref_fig%d"
TAB_BM = "_Ref_tab%s"
ALG_BM = "_Ref_alg%d"
BIB_BM = "_Ref_bib%d"

# Caption forms; group 1 is the token to bookmark.
#
# The trailing period matters. A caption reads "FIGURE 2. Overview of ...",
# whereas a sentence that merely *refers* to the figure reads "Fig. 2 summarises
# ...". Without requiring the period, the lead-in sentence is mistaken for the
# caption, the bookmark lands on the reference instead of the target, and the
# real caption is then skipped as a duplicate.
CAPTIONS = [
    (re.compile(r'^(?:FIGURE|Fig\.)\s*(\d+)\.'), lambda g: FIG_BM % int(g)),
    (re.compile(r'^TABLE\s+([IVXL]+)'),        lambda g: TAB_BM % g),
    (re.compile(r'^Algorithm\s+(\d+)'),        lambda g: ALG_BM % int(g)),
    (re.compile(r'^(\[\d+\])'),                lambda g: BIB_BM % int(g[1:-1])),
]

# In-prose references: (pattern, bookmark namer, group holding the token).
# Group 0 means the whole match becomes the field (bibliography brackets).
REFERENCES = [
    (re.compile(r'\[\d+\]'),             lambda m: BIB_BM % int(m.group(0)[1:-1]), 0),
    (re.compile(r'\bFig\.\s*(\d+)'),     lambda m: FIG_BM % int(m.group(1)),       1),
    (re.compile(r'\bFigure\s+(\d+)'),    lambda m: FIG_BM % int(m.group(1)),       1),
    (re.compile(r'\bTable\s+([IVXL]+)'), lambda m: TAB_BM % m.group(1),            1),
    (re.compile(r'\bAlgorithm\s+(\d+)'), lambda m: ALG_BM % int(m.group(1)),       1),
]


# --------------------------------------------------------------- XML helpers
def _rpr(run_el):
    rpr = run_el.find(qn("w:rPr"))
    return copy.deepcopy(rpr) if rpr is not None else None


def _run(text, rpr, instr=False):
    r = OxmlElement("w:r")
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    t = OxmlElement("w:instrText" if instr else "w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def _fldchar(kind, rpr):
    r = OxmlElement("w:r")
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    f = OxmlElement("w:fldChar")
    f.set(qn("w:fldCharType"), kind)
    r.append(f)
    return r


def ref_field(bookmark, shown, rpr):
    """The five runs Word uses for a REF field with the hyperlink switch."""
    return [_fldchar("begin", rpr),
            _run(" REF %s \\h " % bookmark, rpr, instr=True),
            _fldchar("separate", rpr),
            _run(shown, rpr),
            _fldchar("end", rpr)]


def iter_all_paragraphs(doc):
    """Yield (paragraph, in_table).

    Algorithm floats are rendered as single-cell tables, so "inside a table" is
    what distinguishes the real "Algorithm 2" header from a body sentence that
    merely opens with the words "Algorithm 2".
    """
    for p in doc.paragraphs:
        yield p, False
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p, True


def merge_text_runs(par):
    """Merge adjacent plain-text runs sharing identical formatting.

    Every later edit then happens inside a single run, which removes all
    cross-run offset arithmetic.
    """
    from lxml import etree

    def rpr_key(r):
        el = r.find(qn("w:rPr"))
        return etree.tostring(el) if el is not None else b""

    merged, prev = 0, None
    for r in list(par._p.findall(qn("w:r"))):
        t = r.find(qn("w:t"))
        extra = [c for c in r if c.tag not in (qn("w:rPr"), qn("w:t"))]
        if t is None or extra:                    # image, field char, break
            prev = None
            continue
        if prev is not None and prev.getnext() is r and rpr_key(prev) == rpr_key(r):
            pt = prev.find(qn("w:t"))
            pt.text = (pt.text or "") + (t.text or "")
            pt.set(qn("xml:space"), "preserve")
            par._p.remove(r)
            merged += 1
            continue
        prev = r
    return merged


def split_and_bookmark(par, span_start, span_end, name, bid):
    """Bookmark exactly par.text[span_start:span_end], splitting runs as needed."""
    pos = 0
    for r in list(par._p.findall(qn("w:r"))):
        t = r.find(qn("w:t"))
        if t is None or not t.text:
            continue
        a, b = pos, pos + len(t.text)
        if a <= span_start and span_end <= b:          # wholly inside this run
            rpr, text = _rpr(r), t.text
            # A run may carry siblings of w:t -- most often w:tab, as in the
            # "[1]<tab>" that opens each bibliography entry. Rebuilding from the
            # text alone would silently drop them.
            extras = [copy.deepcopy(c) for c in r
                      if c.tag not in (qn("w:rPr"), qn("w:t"))]
            head = text[:span_start - a]
            mid = text[span_start - a:span_end - a]
            tail = text[span_end - a:]
            nodes = []
            if head:
                nodes.append(_run(head, rpr))
            bstart = OxmlElement("w:bookmarkStart")
            bstart.set(qn("w:id"), str(bid))
            bstart.set(qn("w:name"), name)
            bend = OxmlElement("w:bookmarkEnd")
            bend.set(qn("w:id"), str(bid))
            nodes += [bstart, _run(mid, rpr), bend]
            if tail:
                nodes.append(_run(tail, rpr))
            if extras:
                keep = OxmlElement("w:r")
                if rpr is not None:
                    keep.append(copy.deepcopy(rpr))
                for c in extras:
                    keep.append(c)
                nodes.append(keep)
            anchor = r
            for n in nodes:
                anchor.addnext(n)
                anchor = n
            par._p.remove(r)
            return True
        pos = b
    return False


def pass1_bookmarks(doc):
    bid = 1000
    counts = {"fig": 0, "tab": 0, "alg": 0, "bib": 0}
    seen = set()
    for p, in_table in iter_all_paragraphs(doc):
        txt = p.text
        if not txt.strip():
            continue
        stripped = txt.lstrip()
        lead = len(txt) - len(stripped)
        for rx, namer in CAPTIONS:
            m = rx.match(stripped)
            if not m:
                continue
            if rx.pattern.startswith("^Algorithm") and not in_table:
                continue
            name = namer(m.group(1))
            if name in seen:
                break
            merge_text_runs(p)
            if split_and_bookmark(p, lead + m.start(1), lead + m.end(1), name, bid):
                seen.add(name)
                bid += 1
                counts["fig" if "fig" in name else
                       "tab" if "tab" in name else
                       "alg" if "alg" in name else "bib"] += 1
            break
    return counts, seen


def is_caption(txt, in_table):
    t = txt.lstrip()
    for rx, _ in CAPTIONS:
        if not rx.match(t):
            continue
        if rx.pattern.startswith("^Algorithm") and not in_table:
            continue          # a sentence opening with "Algorithm N", not a header
        return True
    return False


def pass2_fields(doc, known):
    converted, skipped, merged = 0, 0, 0
    for p, in_table in iter_all_paragraphs(doc):
        txt = p.text
        if not txt.strip() or is_caption(txt, in_table):
            continue
        if not any(rx.search(txt) for rx, _, _ in REFERENCES):
            continue
        merged += merge_text_runs(p)

        for r in list(p._p.findall(qn("w:r"))):
            t = r.find(qn("w:t"))
            if t is None or not t.text:
                continue
            text = t.text

            hits = []
            for rx, namer, grp in REFERENCES:
                for m in rx.finditer(text):
                    hits.append((m.start(), m.end(), namer(m),
                                 m.start(grp), m.end(grp)))
            if not hits:
                continue
            hits.sort()
            clean, last = [], -1
            for h in hits:
                if h[0] >= last:
                    clean.append(h)
                    last = h[1]

            rpr = _rpr(r)
            nodes, cur, ok = [], 0, True
            for _s, _e, bm, ts, te in clean:
                if bm not in known:
                    skipped += 1
                    ok = False
                    break
                if text[cur:ts]:                       # e.g. "Fig. " stays literal
                    nodes.append(_run(text[cur:ts], rpr))
                nodes.extend(ref_field(bm, text[ts:te], rpr))
                cur = te
                converted += 1
            if not ok:
                continue
            if text[cur:]:
                nodes.append(_run(text[cur:], rpr))

            anchor = r
            for n in nodes:
                anchor.addnext(n)
                anchor = n
            p._p.remove(r)
    return converted, skipped, merged


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else "Cross_Vehicle_7V_IEEE_Access.docx"
    dst = sys.argv[2] if len(sys.argv) > 2 else src
    doc = docx.Document(src)

    counts, known = pass1_bookmarks(doc)
    print("bookmarks: %d figures, %d tables, %d algorithms, %d references"
          % (counts["fig"], counts["tab"], counts["alg"], counts["bib"]))

    converted, skipped, merged = pass2_fields(doc, known)
    print("fields   : %d converted, %d skipped, %d runs merged"
          % (converted, skipped, merged))

    doc.save(dst)
    print("wrote    : %s" % os.path.basename(dst))
    return 0 if skipped == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
